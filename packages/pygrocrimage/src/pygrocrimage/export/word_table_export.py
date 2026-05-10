import pandas as pd
from docx import Document

from pygrocrimage.export.table_exporter_base import TableExporterBase


class WordTableExporter(TableExporterBase):
    def export(self, df: pd.DataFrame, output_path: str) -> None:

        doc = Document()
        table = doc.add_table(rows=df.shape[0] + 1, cols=df.shape[1])
        table.style = 'Table Grid'

        for j, col in enumerate(df.columns):
            cell = table.cell(0, j)
            cell.text = str(col)
            cell.paragraphs[0].runs[0].bold = True

        for i, (_, row) in enumerate(df.iterrows()):
            for j, val in enumerate(row):
                table.cell(i + 1, j).text = str(val)

        doc.save(output_path)
